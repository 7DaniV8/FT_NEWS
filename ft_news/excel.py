"""
ft_news/excel.py — HISTORICO + RESUMEN de FT NEWS (segunda entrega), con
openpyxl, mismo espíritu que cmp_excel.py de RankingFTR: TODO lo evaluado
(🔥, 🔎 y nada), con sus bandas, nivel, motivo, qué llegó a Discord, el
resultado y la marca "¿era interesante?". El RESUMEN cruza, sin
probabilidades: por tipo de alerta, por nivel, por categoría fuerte, por
banda de mercado, "¿era interesante?" y "¿se movió el mercado hacia la
noticia?" (frecuencias observadas).
"""
from __future__ import annotations

import io
import json
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLS = ("id", "creado_en", "actualizado_en", "tipo_alerta", "motivo", "jugador_x", "rival", "torneo", "circuito",
        "fecha_estimada", "direccion_neta", "nivel", "categoria_fuerte", "categorias", "utr_x", "utr_rival", "delta_utr",
        "cuota_x", "cuota_rival", "banda_utr", "banda_mercado", "desvio", "bandas_calibradas", "conclusion",
        "dicho_discord", "dicho_en", "dicho_motivo", "ganador", "retiro", "score", "cuota_cierre_x", "cuota_cierre_rival",
        "mercado_hacia_noticia", "gano_x", "era_interesante", "marcado_por")


def filas(conn) -> list[dict]:
    out = []
    for r in conn.execute(
            "SELECT e.*, p.torneo, p.fecha_estimada, p.circuito, jx.nombre AS jugador_x, jr.nombre AS rival, "
            "r.ganador_id, r.retiro, r.score, r.cuota_cierre_x, r.cuota_cierre_rival, r.mercado_hacia_noticia, "
            "r.era_interesante, r.marcado_por, jg.nombre AS ganador FROM news_evaluacion e "
            "JOIN news_partido p ON p.id=e.partido_id JOIN news_jugador jx ON jx.id=e.jugador_x_id "
            "JOIN news_jugador jr ON jr.id=e.rival_id LEFT JOIN news_resultado r ON r.evaluacion_id=e.id "
            "LEFT JOIN news_jugador jg ON jg.id=r.ganador_id ORDER BY e.id").fetchall():
        d = dict(r)
        ids = json.loads(d.get("eventos_json") or "[]")
        cats = []
        if ids:
            q = "SELECT categoria FROM news_evento WHERE id IN (%s)" % ",".join("?" * len(ids))
            cats = [x[0] for x in conn.execute(q, ids).fetchall()]
        d["categorias"] = ",".join(sorted(set(cats)))
        d["gano_x"] = None if d.get("ganador_id") is None else int(d["ganador_id"] == d["jugador_x_id"])
        out.append({k: d.get(k) for k in COLS})
    return out


def resumen(fs: list[dict]) -> dict:
    def frec(sub, key):
        c = Counter(x.get(key) for x in sub)
        return {str(k): v for k, v in c.items()}

    def grupo(key):
        g = defaultdict(list)
        for f in fs:
            g[f.get(key)].append(f)
        out = {}
        for k, sub in sorted(g.items(), key=lambda kv: str(kv[0])):
            con_res = [x for x in sub if x.get("gano_x") is not None]
            hacia = [x for x in sub if x.get("mercado_hacia_noticia") is not None]
            out[str(k)] = {"n": len(sub), "dichas": sum(1 for x in sub if x.get("dicho_discord") == 1),
                           "era_interesante": frec(sub, "era_interesante"),
                           "con_resultado": len(con_res),
                           "lado_noticia_acerto": _acierto(con_res),
                           "mercado_hacia_noticia": round(sum(x["mercado_hacia_noticia"] for x in hacia) / len(hacia), 2) if hacia else None}
        return out
    return {"total": len(fs), "por_tipo": grupo("tipo_alerta"), "por_nivel": grupo("nivel"),
            "por_banda_mercado": grupo("banda_mercado"), "por_motivo": grupo("motivo"),
            "por_categoria_fuerte": grupo("categoria_fuerte")}


def _acierto(con_res):
    """Frecuencia con que ganó el lado que la noticia señalaba (NEGATIVA → ganó el rival; POSITIVA → ganó X)."""
    n = ok = 0
    for x in con_res:
        if x.get("direccion_neta") == "NEGATIVA":
            n += 1; ok += 1 - x["gano_x"]
        elif x.get("direccion_neta") == "POSITIVA":
            n += 1; ok += x["gano_x"]
    return round(ok / n, 2) if n else None


def construir(conn) -> bytes:
    fs = filas(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "HISTORICO"
    ws.append(list(COLS))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for f in fs:
        ws.append([f.get(k) for k in COLS])
    for i, k in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, len(k) + 2))
    ws.freeze_panes = "A2"
    r = resumen(fs)
    ws2 = wb.create_sheet("RESUMEN")
    ws2.append(["FT NEWS INTELLIGENCE — RESUMEN (frecuencias observadas, no probabilidades)"])
    ws2["A1"].font = Font(bold=True)
    ws2.append([f"Evaluaciones: {r['total']}"])
    for titulo, clave in (("Por tipo de alerta", "por_tipo"), ("Por nivel de evidencia", "por_nivel"),
                          ("Por banda de mercado", "por_banda_mercado"), ("Por motivo", "por_motivo"),
                          ("Categoría fuerte (1) vs acumulación (0)", "por_categoria_fuerte")):
        ws2.append([])
        ws2.append([titulo])
        ws2[ws2.max_row][0].font = Font(bold=True)
        ws2.append(["grupo", "n", "dichas en Discord", "era interesante SI", "NO", "MERCADO_LO_SABIA", "con resultado",
                    "lado de la noticia acertó (frec.)", "mercado hacia la noticia (frec.)"])
        for k, v in r[clave].items():
            ei = v["era_interesante"]
            ws2.append([k, v["n"], v["dichas"], ei.get("SI", 0), ei.get("NO", 0), ei.get("MERCADO_LO_SABIA", 0),
                        v["con_resultado"], v["lado_noticia_acerto"], v["mercado_hacia_noticia"]])
    ws2.column_dimensions["A"].width = 34
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
